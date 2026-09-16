"""Build a Spanish Word data dictionary from the checked-in public schema."""
from pathlib import Path
import re
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SCHEMA = ROOT / "supabase" / "remote-public-schema.sql"
OUT = ROOT / "docs" / "Diccionario de datos Nodo.docx"

TABLE_EXPLANATIONS = {
    "organizations": "Organizaciones o talleres que operan dentro de Nodo.",
    "profiles": "Perfil y permisos de cada persona usuaria.",
    "customers": "Clientes finales de cada organización.",
    "customer_devices": "Equipos pertenecientes a clientes y recibidos por una organización.",
    "device_types": "Tipos de dispositivo que definen el comportamiento del formulario.",
    "device_receptions": "Recepciones de equipos y su snapshot histórico.",
    "audit_events": "Registro auditable de operaciones relevantes del sistema.",
}

def field_explanation(name: str) -> str:
    if name == "id": return "Identificador único del registro."
    if name in {"created_at", "updated_at"}: return "Fecha y hora de creación o última actualización."
    if name.startswith("fk_"): return "Referencia a otro registro mediante clave foránea."
    if name.endswith("_id"): return "Identificador relacionado con otra entidad."
    if name in {"organization_id", "fk_organizacion_id"}: return "Organización propietaria; permite el aislamiento multiempresa."
    if name in {"is_active", "activo"}: return "Indica si el registro está disponible para uso operativo."
    if name in {"name", "nombre", "label", "title", "titulo"}: return "Nombre o etiqueta visible para las personas usuarias."
    if name in {"code", "key", "clave"}: return "Código técnico estable usado por reglas y relaciones."
    if name in {"description", "descripcion"}: return "Descripción complementaria del registro."
    if name in {"sort_order", "orden"}: return "Orden de presentación dentro de su listado o formulario."
    if name in {"status", "estado"}: return "Estado operativo actual del registro."
    if name.startswith("is_") or name.startswith("has_"): return "Indicador booleano de una condición o característica."
    if "email" in name: return "Correo electrónico asociado al registro."
    if "phone" in name or "telefono" in name: return "Número telefónico de contacto."
    if "metadata" in name or "attributes" in name or "snapshot" in name: return "Datos estructurados adicionales o copia histórica en formato JSON."
    return "Dato propio de esta entidad usado por la operación de Nodo."

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)

def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""; p = cell.paragraphs[0]; run = p.add_run(text); run.bold = bold; run.font.size = Pt(8.5)
    if color: run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def parse_tables(sql: str):
    pattern = re.compile(r'CREATE TABLE IF NOT EXISTS "public"\."([^"]+)" \((.*?)\n\);', re.S)
    tables = []
    for name, body in pattern.findall(sql):
        columns = []
        for raw in body.splitlines():
            line = raw.strip().rstrip(",")
            match = re.match(r'"([^"]+)"\s+(.+)', line)
            if not match or line.upper().startswith(("CONSTRAINT", "PRIMARY", "UNIQUE", "CHECK", "FOREIGN")): continue
            col, definition = match.groups()
            datatype = definition.split(" ", 1)[0].replace('"', '')
            columns.append((col, datatype))
        tables.append((name, columns))
    return tables

def main():
    tables = parse_tables(SCHEMA.read_text(encoding="utf-8"))
    doc = Document(); section = doc.sections[0]; section.top_margin = Inches(.65); section.bottom_margin = Inches(.65); section.left_margin = Inches(.65); section.right_margin = Inches(.65)
    normal = doc.styles["Normal"]; normal.font.name = "Aptos"; normal.font.size = Pt(9)
    for style in ["Title", "Heading 1", "Heading 2"]: doc.styles[style].font.name = "Aptos"; doc.styles[style].font.color.rgb = RGBColor(0,0,0)
    title = doc.add_paragraph(style="Title"); title.alignment = WD_ALIGN_PARAGRAPH.CENTER; title.add_run("Diccionario de datos Nodo")
    subtitle = doc.add_paragraph(); subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER; subtitle.add_run("Guía en español de las tablas y campos del esquema público").italic = True
    doc.add_paragraph(f"Este documento describe las {len(tables)} tablas detectadas en el esquema público exportado del proyecto. Cada tabla incluye su finalidad y el significado práctico de sus campos. Las claves foráneas relacionan datos; los campos de organización sostienen el aislamiento multiempresa.")
    doc.add_heading("Cómo leer esta guía", level=1)
    doc.add_paragraph("Los campos llamados id identifican un registro. Los terminados en id o iniciados por fk enlazan con otra tabla. Los campos is_active o activo permiten desactivar registros sin borrar historial. Los campos JSON almacenan información estructurada complementaria.")
    for table_name, columns in tables:
        doc.add_page_break()
        doc.add_heading(table_name, level=1)
        doc.add_paragraph(TABLE_EXPLANATIONS.get(table_name, "Tabla del esquema público de Nodo. Su finalidad concreta se deriva de sus relaciones y campos operativos."))
        doc.add_paragraph(f"Campos documentados: {len(columns)}.")
        table = doc.add_table(rows=1, cols=3); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.style = "Table Grid"
        headers = ["Campo", "Tipo", "Uso en Nodo"]
        for cell, text in zip(table.rows[0].cells, headers): shade(cell, "163A5F"); set_cell_text(cell, text, True, (255,255,255))
        for index, (name, datatype) in enumerate(columns):
            cells = table.add_row().cells
            if index % 2 == 1:
                for cell in cells: shade(cell, "EEF4F8")
            set_cell_text(cells[0], name, True); set_cell_text(cells[1], datatype); set_cell_text(cells[2], field_explanation(name))
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER; footer.add_run("Nodo  |  Diccionario de datos")
    OUT.parent.mkdir(exist_ok=True); doc.save(OUT); print(f"{OUT}\nTables: {len(tables)}")

if __name__ == "__main__": main()
